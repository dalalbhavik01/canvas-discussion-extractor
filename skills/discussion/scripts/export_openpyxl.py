"""Portable Excel exporter for hosts without Codex's Artifact Tool."""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export(directory):
    directory = Path(directory)
    prepared = json.loads((directory / ".audit" / "prepared.json").read_text(encoding="utf-8"))
    for section in prepared["sections"]:
        spec = section["sheets"][0]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = spec["name"]
        for row_number, values in enumerate(spec["values"], 1):
            for column_number, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column_number, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number == 1:
                    cell.fill = PatternFill("solid", fgColor="4A1730")
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                else:
                    cell.font = Font(name="Arial", size=10, color="202020")
        sheet.column_dimensions["A"].width = 28
        for column in range(2, len(spec["values"][0]) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 70
        sheet.row_dimensions[1].height = 32
        for row in range(2, len(spec["values"]) + 1):
            sheet.row_dimensions[row].height = 100
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        workbook.save(directory / (section["key"] + ".xlsx"))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python export_openpyxl.py OUTPUT_DIRECTORY")
    export(sys.argv[1])
